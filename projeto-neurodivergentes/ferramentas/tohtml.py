from openpyxl import load_workbook
import html
P="/home/user/readme/projeto-neurodivergentes/"
wb=load_workbook(P+"agenda-atendimentos.xlsx",data_only=True)
order=["Resumo","Comparativo","Calendário","Grade do dia","Blocos","Semanal 16 sessões","Médicos","Agenda médica","Serviços extras","Equipe e turnos"]
def fmt(v,pct=False):
    if pct and isinstance(v,(int,float)): return f"{round(v*100)}%"
    if v is None: return ""
    if isinstance(v,float):
        if abs(v-round(v))<1e-9: return f"{int(round(v)):,}".replace(",",".")
        return f"{v:.1f}".replace(".",",")
    if isinstance(v,int): return f"{v:,}".replace(",",".")
    return html.escape(str(v))
secs=[]
for n in order:
    ws=wb[n]; rows=[]
    merged={}
    for mr in ws.merged_cells.ranges:
        merged[(mr.min_row,mr.min_col)]=(mr.max_row-mr.min_row+1,mr.max_col-mr.min_col+1)
        for r in range(mr.min_row,mr.max_row+1):
            for c in range(mr.min_col,mr.max_col+1):
                if (r,c)!=(mr.min_row,mr.min_col): merged[(r,c)]=None
    out=[]
    for r in range(1,ws.max_row+1):
        cells=[]; has=False
        nn=[c for c in range(1,ws.max_column+1) if ws.cell(r,c).value is not None]
        if len(nn)==1 and nn[0]==1 and ws.cell(r,1).fill.fill_type is None and (r,1) not in merged:
            b=' class=b' if ws.cell(r,1).font.bold else ''
            out.append(f'<tr><td{b} colspan="{ws.max_column}" style="border:0;padding-left:0">{fmt(ws.cell(r,1).value)}</td></tr>'); continue
        for c in range(1,ws.max_column+1):
            if (r,c) in merged and merged[(r,c)] is None: continue
            cell=ws.cell(r,c); v=cell.value
            if v is not None: has=True
            fill=cell.fill.fgColor.rgb if cell.fill and cell.fill.fill_type=="solid" else None
            style=""; tag="td"
            if fill and fill not in ("00000000",):
                col="#"+fill[-6:]
                if col.upper()=="#1F4E79": tag="th"
                else: style=f' style="background:{col};color:#1a1a1a"'
            if v is None and not fill: cells.append("<td class=e></td>"); continue
            span=""
            if (r,c) in merged and merged[(r,c)]:
                rs,cs=merged[(r,c)]; span=f' colspan="{cs}"'+(f' rowspan="{rs}"' if rs>1 else "")
            cls=" class=b" if (cell.font and cell.font.bold and tag=="td") else ""
            cells.append(f"<{tag}{style}{cls}{span}>{fmt(v,cell.number_format=='0%').replace(chr(10),'<br>')}</{tag}>")
        if has: out.append("<tr>"+"".join(cells)+"</tr>")
    secs.append((n,"".join(out)))
tabs="".join(f'<button data-t="{i}"{" class=on" if i==0 else ""}>{html.escape(n)}</button>' for i,(n,_) in enumerate(secs))
body="".join(f'<section id="s{i}"{"" if i==0 else " hidden"}><div class=w><table>{t}</table></div></section>' for i,(n,t) in enumerate(secs))
page=f'''<!doctype html><html lang="pt-BR"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>Agenda dos atendimentos</title>
<style>
:root{{--bg:#fff;--fg:#1a1a1a;--mut:#555;--line:#c8ccd0;--acc:#1f4e79}}
@media (prefers-color-scheme:dark){{:root{{--bg:#16181b;--fg:#e8e8e8;--mut:#a0a4a8;--line:#3a3f45;--acc:#6aa7d8}}}}
body{{margin:0;font:14px/1.45 Arial,Helvetica,sans-serif;background:var(--bg);color:var(--fg)}}
header{{padding:16px}}h1{{font-size:20px;margin:0 0 4px}}p{{margin:0;color:var(--mut)}}
nav{{display:flex;gap:6px;flex-wrap:wrap;padding:0 16px 12px;position:sticky;top:0;background:var(--bg)}}
nav button{{font:inherit;padding:6px 12px;border:1px solid var(--line);background:none;color:var(--fg);border-radius:6px;cursor:pointer}}
nav button.on{{background:var(--acc);color:#fff;border-color:var(--acc)}}
.w{{overflow:auto;padding:0 16px 24px}}table{{border-collapse:collapse}}
td,th{{border:1px solid var(--line);padding:4px 8px;text-align:left;vertical-align:middle}}
th{{background:#1f4e79;color:#fff;text-align:center;position:sticky;top:0}}td.b{{font-weight:bold}}td.e{{border:0}}
</style></head><body>
<header><h1>Instituto Affectus – agenda dos atendimentos</h1><p>Plano adotado: 16 sessões por mês, 1 visita por semana (julho/2026 a junho/2027). Os mesmos dados estão na planilha agenda-atendimentos.xlsx.</p></header>
<nav>{tabs}</nav>{body}
<script>document.querySelectorAll("nav button").forEach(b=>b.onclick=()=>{{document.querySelectorAll("nav button").forEach(x=>x.classList.toggle("on",x===b));document.querySelectorAll("section").forEach((s,i)=>s.hidden=i!=b.dataset.t)}});const h=parseInt(location.hash.slice(1));if(h>=0)document.querySelectorAll("nav button")[h]?.click()</script>
</body></html>'''
open(P+"agenda-atendimentos.html","w").write(page)
print(len(page))
