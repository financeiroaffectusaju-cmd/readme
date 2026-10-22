import json, sys, html
from datetime import date, timedelta
from collections import Counter, defaultdict
sys.path.insert(0,'.')
from extras import run as run_extras
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L

OUT="/home/user/readme/projeto-neurodivergentes/"
B=json.load(open("blocks.json"))
ROOMS=[("Psi","Psicologia",2),("TO","Terapia Ocupacional",2),("Fono","Fonoaudiologia",2),("Psped","Psicopedagogia",2),
       ("Psmot","Psicomotricidade",2),("Music","Musicoterapia",1),("Alim","Terapia alimentar",2)]
NAME={k:n for k,n,_ in ROOMS}
PKG={1:["Psi","TO","Psped","Psmot"],2:["Psi","Fono","Psped","Alim"],3:["TO","Fono","Psped","Music"],
     4:["Psi","TO","Psmot","Alim"],5:["Psi","Fono","Psmot","Music"],6:["TO","Fono","Psmot","Alim"]}
PCOL={1:"BDD7EE",2:"C6E0B4",3:"FFE699",4:"F8CBAD",5:"D9C3E8",6:"B7E1E0"}
def tm(s):
    m=7*60+30*s if s<10 else 13*60+30*(s-10)
    return f"{m//60:02d}:{m%60:02d}"
def tend(s):
    m=(7*60+30*s if s<10 else 13*60+30*(s-10))+30
    return f"{m//60:02d}:{m%60:02d}"
DAYS=["Seg S1","Ter S1","Qua S1","Qui S1","Sex S1","Seg S2","Ter S2","Qua S2","Qui S2","Sex S2"]

# ---------- blocos e crianças ----------
PATS={1:[0,1,2,3],2:[4,5,6,7],3:[8,9,0],4:[1,2,3],5:[4,5,6],6:[7,8,9]}
FREE={}
for pk,pl in PATS.items():
    days=sorted(d for p in pl for d in (p%10,(p+3)%10,(p+6)%10))
    idx=[i for i,b in enumerate(B) if b["pkg"]==pk]
    assert len(idx)==len(days),(pk,len(idx),len(days))
    for k,i in enumerate(idx): FREE[i]=days[(k*7)%len(days)]
blocks=[]
for i,b in enumerate(B):
    bid=f"B{i+1:02d}"; o=(FREE[i]+1)%10
    kids=[]
    for j,kn in enumerate("ABC"):
        kid=f"C{i*3+j+1:03d}"
        days=sorted([(o+j)%10,(o+j+3)%10,(o+j+6)%10])
        kids.append((kid,days))
    free=(o+9)%10
    s0=b["seq"][0][0]
    blocks.append(dict(id=bid,pkg=b["pkg"],seq=b["seq"],start=tm(s0),end=tend(s0+3),kids=kids,free=free))
# grade do dia: salas x slots
SEATCOLS=[(r,s) for r,_,n in ROOMS for s in range(1,n+1)]
grid={}
for sl in range(22):
    used=defaultdict(int)
    for blk in blocks:
        for t,r in blk["seq"]:
            if t==sl:
                used[r]+=1; grid[(sl,r,used[r])]=blk
# ---------- calendário ----------
meses=[("Jul/2026",22),("Ago/2026",21),("Set/2026",21),("Out/2026",21),("Nov/2026",19),("Dez/2026",21),
       ("Jan/2027",20),("Fev/2027",18),("Mar/2027",21),("Abr/2027",21),("Mai/2027",20),("Jun/2027",21)]
fer=[("08/07/2026","Qua","Emancipação Política de Sergipe (estadual)"),("07/09/2026","Seg","Independência (nacional)"),
 ("12/10/2026","Seg","N. Sra. Aparecida (nacional)"),("02/11/2026","Seg","Finados (nacional)"),("20/11/2026","Sex","Consciência Negra (nacional)"),
 ("08/12/2026","Ter","N. Sra. da Conceição (municipal Aracaju) – conferir decreto"),("25/12/2026","Sex","Natal (nacional)"),
 ("01/01/2027","Sex","Confraternização Universal (nacional)"),("08/02/2027","Seg","Carnaval (ponto facultativo)"),
 ("09/02/2027","Ter","Carnaval (ponto facultativo)"),("17/03/2027","Qua","Aniversário de Aracaju (municipal)"),
 ("26/03/2027","Sex","Sexta-feira Santa (nacional)"),("21/04/2027","Qua","Tiradentes (nacional)"),
 ("27/05/2027","Qui","Corpus Christi (confirmado: sem atendimento)"),("24/06/2027","Qui","São João (municipal Aracaju) – EM ABERTO: ver célula B19")]
hol=set()
for d,_,_ in fer:
    dd,mm,yy=map(int,d.split("/")); hol.add(date(yy,mm,dd))
# médicos: sugestão de datas
def nth_wd(y,m,wd,n):
    d=date(y,m,1); c=0
    while True:
        if d.weekday()==wd:
            c+=1
            if c==n: return d
        d+=timedelta(1)
def free_day(d):
    while d in hol or d.weekday()>4: d+=timedelta(1)
    return d
ml=[(2026,7),(2026,8),(2026,9),(2026,10),(2026,11),(2026,12),(2027,1),(2027,2),(2027,3),(2027,4),(2027,5),(2027,6)]
med=[]
for y,m in ml:
    med.append((f"{m:02d}/{y}",free_day(nth_wd(y,m,1,2)),free_day(nth_wd(y,m,2,3)),free_day(nth_wd(y,m,3,4))))
fmt=lambda d: d.strftime("%d/%m (")+["seg","ter","qua","qui","sex"][d.weekday()]+")"
# extras
vs,asg=run_extras(60)
cnt=Counter((t,h) for t,h in asg.values())
HRS=[7,8,9,10,11,13,14,15,16,17,18]
# oficinas
OF=["Confeitaria e panificação (mães)","Artesanato e personalizados (mães)","Estética básica (mães)",
    "Elétrica e instalações básicas (pais)","Informática e manutenção de celulares (pais)","Barbearia (pais)"]
OSL=[("07:00","08:30"),("08:30","10:00"),("10:00","11:30"),("13:00","14:30"),("14:30","16:00"),("16:00","17:30"),("17:30","19:00")]
ogrid=[[OF[(d*7+j)%6] for d in range(5)] for j in range(7)]
# psicóloga dos pais
GRP={"08:30","10:30","14:30","16:30"}
# dentista 40 min
def dent():
    out=[]
    for start,n in ((7*60,7),(13*60,9)):
        for k in range(n):
            a=start+40*k; b=a+40
            out.append(f"{a//60:02d}:{a%60:02d}–{b//60:02d}:{b%60:02d}")
    return out
DENT=dent()
# equipe
EQ=[("Fonoaudiologia",2,2),("Terapia Ocupacional",2,2),("Psicologia",2,2),("Psicopedagogia",2,2),("Psicomotricidade",2,2),
    ("Musicoterapia",1,2),("Terapia alimentar (nutrição)",1,2),("Psicóloga dos pais",1,2),("Odontologia (1 por dia)",1,2),
    ("Capoeira (1 por vez)",1,2),("Xadrez (1 por turno)",1,2)]
json.dump(dict(blocks=blocks),open("data.json","w"))

# =============== XLSX ===============
wb=Workbook()
F="Arial"
hf=Font(name=F,bold=True,color="FFFFFF"); hfill=PatternFill("solid",fgColor="1F4E79")
bf=Font(name=F); bold=Font(name=F,bold=True); blue=Font(name=F,color="0000FF"); title=Font(name=F,bold=True,size=14)
thin=Side(style="thin",color="BFBFBF"); bd=Border(left=thin,right=thin,top=thin,bottom=thin)
ctr=Alignment(horizontal="center",vertical="center",wrap_text=True)
def hdr(ws,r,vals,c0=1):
    for i,v in enumerate(vals):
        c=ws.cell(r,c0+i,v); c.font=hf; c.fill=hfill; c.alignment=ctr; c.border=bd
def put(ws,r,c,v,font=bf,fill=None,al=None,b=True):
    x=ws.cell(r,c,v); x.font=font
    if fill: x.fill=PatternFill("solid",fgColor=fill)
    if al: x.alignment=al
    if b: x.border=bd
    return x

# --- Calendário (primeiro, para referências)
wc=wb.active; wc.title="Calendário"
wc["A1"]="Calendário de funcionamento – jul/2026 a jun/2027 (Aracaju/SE)"; wc["A1"].font=title
hdr(wc,3,["Mês","Dias de funcionamento"])
for i,(m,n) in enumerate(meses):
    put(wc,4+i,1,m); put(wc,4+i,2,n,blue)
put(wc,16,1,"Total no ano",bold); put(wc,16,2,"=SUM(B4:B15)",bold)
wc["D3"]="Feriados descontados"; wc["D3"].font=bold
hdr(wc,4,["Data","Dia","Feriado"],4)
for i,(d,w,n) in enumerate(fer):
    put(wc,5+i,4,d); put(wc,5+i,5,w); put(wc,5+i,6,n)
wc["B15"]='=21+IF($B$19="S",1,0)'
put(wc,19,1,"São João (24/06/2027) funciona? (S/N)",bold); put(wc,19,2,"N",blue,"FFFF00",ctr)
wc["A20"]="Decisão em aberto: com 'N' o dia é descontado (246 dias); com 'S' o total passa para 247."; wc["A20"].font=bf

from datetime import date as _d, timedelta as _td
_occ=[0]*5; _x=_d(2026,7,1)
while _x<=_d(2027,6,30):
    if _x.weekday()<5: _occ[_x.weekday()]+=1
    _x+=_td(1)
_hc=[0]*5
for _h in hol:
    if _d(2026,7,1)<=_h<=_d(2027,6,30) and _h.weekday()<5: _hc[_h.weekday()]+=1
hdr(wc,3,["Dia da semana","Ocorrências no período","Feriados","Visitas por criança","Sessões por criança"],8)
for _i,_n in enumerate(["Segunda","Terça","Quarta","Quinta","Sexta"]):
    _r=4+_i; put(wc,_r,8,_n); put(wc,_r,9,_occ[_i],blue); put(wc,_r,10,_hc[_i],blue); put(wc,_r,11,f"=I{_r}-J{_r}"); put(wc,_r,12,f"=K{_r}*Resumo!$B$9")
wc["H10"]="Como cada criança vem 1 vez por semana no mesmo dia, o número de feriados do seu dia reduz as visitas do ano. Inclui São João (24/06) e 8/12."; wc["H10"].font=bf
for _c,_w in zip("HIJKL",(16,14,10,14,14)): wc.column_dimensions[_c].width=_w
wc.row_dimensions[3].height=32
wc["A18"]="Base: 261 dias úteis seg–sex no período, menos 15 feriados em dia útil. Itens marcados 'conferir decreto' dependem do decreto oficial da Prefeitura."; wc["A18"].font=bf
for col,w in zip("ABCDEF",(16,22,3,14,8,70)): wc.column_dimensions[col].width=w

# --- Resumo
ws=wb.create_sheet("Resumo",0)
ws["A1"]="Instituto Affectus – plano adotado: 16 sessões por mês, 1 visita de 2 h por semana"; ws["A1"].font=title
ws["A2"]="Decisões da gestora em 22/10/2026: Opção 2 (16 sessões, 1 visita por semana) e Cenário 1 (300 crianças). Azul = premissa editável; o restante é calculado."; ws["A2"].font=bf
hdr(ws,4,["Premissa","Valor","Observação"])
rows=[("Dias de atendimento por ano","=Calendário!B16","","Ver aba Calendário"),
("Horário de funcionamento","07h–19h","","Seg–sex; almoço 12h–13h"),
("Formato das sessões","4 × 30 min","","Todas de 30 min, horários seguidos, 2 h por visita"),
("Sessões de terapia por criança a cada 4 semanas",16,"i","Decisão de 22/10/2026 (antes 24); xadrez, capoeira, oficinas, médicos e odontologia são extras"),
("Sessões por visita",4,"i","3 ou 4 especialidades; modelo usa 4"),
("Visitas por criança a cada 4 semanas","=B8/B9","","1 visita por semana, sempre no mesmo dia, horário, bloco e sala"),
("Dias úteis por semana",5,"i","Cada bloco recebe 1 criança por dia da semana"),
("Blocos de visita por dia",60,"i","Comprovado por solver: 60 visitas/dia cabem com horários seguidos e vagas das salas"),
("Horas de terapia por criança a cada 4 semanas","=B8*0.5","","16 sessões × 30 min")]
for i,(a_,b_,t,c_) in enumerate(rows):
    r=5+i; put(ws,r,1,a_); put(ws,r,2,b_,blue if t=="i" else bf); put(ws,r,3,c_)
ws["A15"]="Cenários"; ws["A15"].font=bold
hdr(ws,16,["Indicador","Cenário 1: 100% de ocupação (ADOTADO)","Cenário 2: 90% (10% reposição), alternativa"])
put(ws,17,1,"Blocos livres por dia (reposição)"); put(ws,17,2,0,blue); put(ws,17,3,6,blue)
put(ws,18,1,"Visitas por dia"); put(ws,18,2,"=$B$12-B17"); put(ws,18,3,"=$B$12-C17")
put(ws,19,1,"Crianças com pacote completo"); put(ws,19,2,"=B18*$B$11",bold); put(ws,19,3,"=C18*$B$11",bold)
put(ws,20,1,"Ocupação dos blocos"); put(ws,20,2,"=B18/$B$12"); put(ws,20,3,"=C18/$B$12")
put(ws,21,1,"Sessões de 30 min por dia"); put(ws,21,2,"=B18*$B$9"); put(ws,21,3,"=C18*$B$9")
put(ws,22,1,"Sessões de terapia por ano da clínica"); put(ws,22,2,"=B19*$B$8*$B$5/20"); put(ws,22,3,"=C19*$B$8*$B$5/20")
put(ws,23,1,"Sessões de terapia por criança por ano"); put(ws,23,2,"=$B$8*$B$5/20"); put(ws,23,3,"=$B$8*$B$5/20")
put(ws,24,1,"Crianças que consultam cada médico no 1º ano (192 por médico/ano)"); put(ws,24,2,"=MIN(1,192/B19)"); put(ws,24,3,"=MIN(1,192/C19)")
put(ws,25,1,"Horário da criança"); put(ws,25,2,"Fixo: mesmo dia, horário, sala e ordem"); put(ws,25,3,"Fixo: mesmo dia, horário, sala e ordem")
for c in "BC":
    ws[c+"20"].number_format="0%"; ws[c+"22"].number_format="#,##0"; ws[c+"23"].number_format="0"; ws[c+"24"].number_format="0%"
for r in range(17,26):
    for c in (2,3): ws.cell(r,c).alignment=ctr
ws["A27"]="Notas"; ws["A27"].font=bold
notes=["Cada bloco de 2 h recebe 1 criança por dia da semana (5 por semana); a criança volta sempre no mesmo dia, horário e sala. Grade nominal na aba Semanal 16 sessões.",
"Cenário 1 adotado: as 300 vagas estão ocupadas e não há vaga reservada para reposição. A vaga de quem falta no dia pode repor outra criança do mesmo pacote. No Cenário 2 (270) as vagas C271–C300 ficariam livres.",
"Teste de encaixe: 60 visitas por dia comprovadas pelo solver; mais de 62 não foram provadas nem descartadas. Tetos teóricos calculados antes da grade não foram obtidos na grade real.",
"Pacotes-modelo (aba Grade do dia) são ilustrativos; a composição real depende da triagem. Todos têm pelo menos 2 entre Psicologia, TO e Fono.",
"Feriado que cai no dia da semana de uma criança faz ela perder a visita daquela semana; usar as vagas livres para repor.",
"Consulta médica (1 dia por mês por médico): com o rodízio da aba Médicos, todas as 300 crianças têm pelo menos 1 consulta no primeiro ano e o ciclo completo com os 3 médicos fecha em 20 meses.",
"As opções avaliadas e descartadas (24 sessões a cada 2 semanas e 24 sessões semanais de 3 h) estão na aba Comparativo."]
for i,n in enumerate(notes): ws.cell(28+i,1,n).font=bf
for col,w in zip("ABC",(62,34,70)): ws.column_dimensions[col].width=w

# --- Grade do dia
wg=wb.create_sheet("Grade do dia")
wg["A1"]="Grade do dia (igual de segunda a sexta) – 60 blocos de visita"; wg["A1"].font=title
wg["A2"]="Cada célula mostra o bloco (B01…B60) que ocupa aquela vaga naquele horário. Cor = pacote-modelo."; wg["A2"].font=bf
hdr(wg,4,["Horário"]+[f"{NAME[r]}\nvaga {s}" if dict((k,n) for k,_,n in ROOMS)[r]>1 else NAME[r] for r,s in SEATCOLS])
rr=5
for sl in range(22):
    if sl==10:
        for c in range(1,len(SEATCOLS)+2): put(wg,rr,c,"Almoço" if c==1 else "",bold,"D9D9D9",ctr)
        rr+=1
    put(wg,rr,1,f"{tm(sl)}–{tend(sl)}",bold,None,ctr)
    for ci,(r,s) in enumerate(SEATCOLS):
        blk=grid.get((sl,r,s))
        if blk: put(wg,rr,2+ci,blk["id"],bf,PCOL[blk["pkg"]],ctr)
        else: put(wg,rr,2+ci,"",bf,None,ctr)
    rr+=1
rr+=1
wg.cell(rr,1,"Pacotes-modelo").font=bold; rr+=1
for k,rl in PKG.items():
    n=sum(1 for b in blocks if b["pkg"]==k)
    put(wg,rr,1,f"Pacote {k}",bold,PCOL[k]); put(wg,rr,2,", ".join(NAME[x] for x in rl)); put(wg,rr,6,f"{n} blocos/dia",bf)
    wg.merge_cells(start_row=rr,start_column=2,end_row=rr,end_column=5); rr+=1
wg.column_dimensions["A"].width=14
for i in range(len(SEATCOLS)): wg.column_dimensions[L(2+i)].width=14
wg.row_dimensions[4].height=44
wg.freeze_panes="B5"

# --- Blocos
wbk=wb.create_sheet("Blocos")
wbk["A1"]="Blocos de visita: horários e salas"; wbk["A1"].font=title
hdr(wbk,3,["Bloco","Pacote","Início","Fim","1ª sessão","2ª sessão","3ª sessão","4ª sessão"])
for i,b in enumerate(blocks):
    r=4+i; put(wbk,r,1,b["id"],bold,PCOL[b["pkg"]]); put(wbk,r,2,b["pkg"]); put(wbk,r,3,b["start"]); put(wbk,r,4,b["end"])
    for j,(t,room) in enumerate(b["seq"]): put(wbk,r,5+j,f"{tm(t)} {NAME[room]}")
for c,w in zip("ABCDEFGH",(9,8,9,9,28,28,28,28)): wbk.column_dimensions[c].width=w
wbk.freeze_panes="A4"

# --- Ciclo
wy=wb.create_sheet("Ciclo Cenário 2")
wy["A1"]="Cenário 2: quem ocupa cada bloco em cada dia do ciclo (10 dias úteis)"; wy["A1"].font=title
wy["A2"]="Cada criança (C001–C180) vem 3 vezes por ciclo, sempre no mesmo bloco. 'livre' = vaga de reposição daquele horário. Trocar C### pelo nome real."; wy["A2"].font=bf
hdr(wy,4,["Bloco","Início","Pacote"]+DAYS)
for i,b in enumerate(blocks):
    r=5+i; put(wy,r,1,b["id"],bold,PCOL[b["pkg"]]); put(wy,r,2,b["start"]); put(wy,r,3,b["pkg"])
    for d in range(10):
        v="livre"
        for kid,days in b["kids"]:
            if d in days: v=kid
        put(wy,r,4+d,v,bf,"EDEDED" if v=="livre" else None,ctr)
for c,w in zip("ABC",(9,9,8)): wy.column_dimensions[c].width=w
for i in range(10): wy.column_dimensions[L(4+i)].width=10
wy.freeze_panes="D5"


# --- Crianças flexíveis (Cenário 1)
flex=[]; fid=0
avail={(pk,d):[i for i,b in enumerate(blocks) if b["pkg"]==pk and b["free"]==d] for pk in PATS for d in range(10)}
for pk,pl in PATS.items():
    for p in pl:
        fid+=1; vl=[]
        for d in sorted((p%10,(p+3)%10,(p+6)%10)):
            vl.append((avail[(pk,d)].pop(0),d))
        flex.append((f"F{fid:02d}",pk,vl))
assert all(not v for v in avail.values())
flexcell={}
for f,pk,vl in flex:
    for bi,d in vl: flexcell[(bi,d)]=f
assert len(flex)==20 and len(flexcell)==60

wy1=wb.create_sheet("Ciclo Cenário 1")
wy1["A1"]="Cenário 1: ocupação de 100%. Quem ocupa cada bloco em cada dia do ciclo (10 dias úteis)"; wy1["A1"].font=title
wy1["A2"]="C001–C180: crianças fixas (mesmo horário sempre). F01–F20: crianças flexíveis, que ocupam as vagas que no Cenário 2 ficam livres para reposição."; wy1["A2"].font=bf
hdr(wy1,4,["Bloco","Início","Pacote"]+DAYS)
for i,b in enumerate(blocks):
    r=5+i; put(wy1,r,1,b["id"],bold,PCOL[b["pkg"]]); put(wy1,r,2,b["start"]); put(wy1,r,3,b["pkg"])
    for d in range(10):
        v=flexcell.get((i,d))
        if v: put(wy1,r,4+d,v,bold,"FFF2CC",ctr)
        else:
            for kid,days in b["kids"]:
                if d in days: v=kid
            put(wy1,r,4+d,v,bf,None,ctr)
for c,w in zip("ABC",(9,9,8)): wy1.column_dimensions[c].width=w
for i in range(10): wy1.column_dimensions[L(4+i)].width=10
wy1.freeze_panes="D5"

wk=wb.create_sheet("Crianças")
wk["A1"]="Lista de vagas de criança (substituir o código pelo nome)"; wk["A1"].font=title
hdr(wk,3,["Código","Tipo","Cenários","Pacote","Dias do ciclo","Bloco e horário"])
r=4
for i,b in enumerate(blocks):
    for kid,days in b["kids"]:
        put(wk,r,1,kid,bold,PCOL[b["pkg"]]); put(wk,r,2,"Fixa"); put(wk,r,3,"1 e 2"); put(wk,r,4,b["pkg"])
        put(wk,r,5,", ".join(DAYS[d] for d in days)); put(wk,r,6,f"{b['id']}, {b['start']}–{b['end']}"); r+=1
for f,pk,vl in flex:
    put(wk,r,1,f,bold,PCOL[pk]); put(wk,r,2,"Flexível"); put(wk,r,3,"1"); put(wk,r,4,pk)
    put(wk,r,5,", ".join(DAYS[d] for bi,d in vl)); put(wk,r,6,"; ".join(f"{DAYS[d]} {blocks[bi]['id']} {blocks[bi]['start']}–{blocks[bi]['end']}" for bi,d in vl)); r+=1
for c,w in zip("ABCDEF",(10,10,10,8,38,70)): wk.column_dimensions[c].width=w
wk.freeze_panes="A4"


# --- Semanal 16 sessões
WD=["Seg","Ter","Qua","Qui","Sex"]
bypk=defaultdict(list)
for i,b in enumerate(blocks): bypk[b["pkg"]].append(i)
reserva=set()
for d in range(5):
    for pk,lst in bypk.items(): reserva.add((lst[(2*d+pk)%len(lst)],d))
assert len(reserva)==30
ws2=wb.create_sheet("Semanal 16 sessões")
ws2["A1"]="Plano adotado: 16 sessões por mês, 1 visita de 2 h por semana. Quem ocupa cada bloco em cada dia da semana"; ws2["A1"].font=title
ws2["A2"]="Cada criança tem um dia da semana, um bloco e um horário fixos (sempre os mesmos). Células amarelas = as 30 vagas que ficam livres para reposição no cenário de 90% (270 crianças); no de 100% (300 crianças) elas são ocupadas por C271–C300."; ws2["A2"].font=bf
hdr(ws2,4,["Bloco","Início","Pacote"]+WD)
nxt=1; resid=[]
for i,b in enumerate(blocks):
    for d in range(5):
        if (i,d) not in reserva: resid.append((i,d))
ids={}
for n,(i,d) in enumerate(resid): ids[(i,d)]=f"C{n+1:03d}"
for n,(i,d) in enumerate(sorted(reserva)): ids[(i,d)]=f"C{271+n:03d}"
for i,b in enumerate(blocks):
    r=5+i; put(ws2,r,1,b["id"],bold,PCOL[b["pkg"]]); put(ws2,r,2,b["start"]); put(ws2,r,3,b["pkg"])
    for d in range(5):
        put(ws2,r,4+d,ids[(i,d)],bf,"FFF2CC" if (i,d) in reserva else None,ctr)
for c,w in zip("ABC",(9,9,8)): ws2.column_dimensions[c].width=w
for i in range(5): ws2.column_dimensions[L(4+i)].width=11
ws2.freeze_panes="D5"


# --- Médicos
import sys as _sys; _sys.path.insert(0,'.')
from medical import plan as _plan, DOCS as _DOCS, doctor_dates as _dd, hm as _hm
from collections import Counter as _C
_rows,_seen,_cnt,_kt,_byw=_plan(ids,blocks,hol,30)
_dist=_C(_cnt.values())
WDN=["seg","ter","qua","qui","sex"]
wm=wb.create_sheet("Médicos")
wm["A1"]="Atendimento médico: 1 dia por mês por médico, com rodízio dos dias da semana"; wm["A1"].font=title
wm["A2"]="Disponibilidade dos médicos: 1 dia por mês cada (8h–12h e 13h–17h). Azul = premissa editável; resultados de distribuição vêm de simulação com a grade semanal."; wm["A2"].font=bf
hdr(wm,4,["Indicador","Valor"])
_r=[("Duração da consulta (min)",30,"i"),("Consultas por dia médico","=480/B5",""),("Dias médicos por ano (3 médicos × 12 meses)",36,"i"),
("Consultas por ano","=B6*B7",""),("Crianças",300,"i"),("Consultas por criança por ano (média)","=B8/B9",""),
("Consultas para cada criança passar pelos 3 médicos","=B9*3",""),("Meses até todas passarem pelos 3 médicos (simulado)",20,"i")]
for i,(a_,b_,t) in enumerate(_r):
    put(wm,5+i,1,a_); put(wm,5+i,2,b_,blue if t=="i" else bf,None,ctr)
wm["B10"].number_format="0.00"
wm["A14"]="Resultado no primeiro ano (30 min)"; wm["A14"].font=bold
hdr(wm,15,["Consultas por criança","Crianças","% das crianças"])
for i,k in enumerate((1,2,3)):
    put(wm,16+i,1,f"{k} consulta(s), com médicos diferentes",bf); put(wm,16+i,2,_dist.get(k,0),blue,None,ctr); put(wm,16+i,3,f"=B{16+i}/$B$9",bf,None,ctr); wm[f"C{16+i}"].number_format="0%"
put(wm,19,1,"Crianças sem nenhuma consulta",bf); put(wm,19,2,"=B9-SUM(B16:B18)",bf,None,ctr); put(wm,19,3,"=B19/$B$9",bf,None,ctr); wm["C19"].number_format="0%"
wm["A21"]="Se a consulta for mais curta (a conversar com os médicos)"; wm["A21"].font=bold
hdr(wm,22,["Duração (min)","Consultas por dia médico","Consultas por ano","Meses até todas passarem pelos 3 médicos (simulado)"])
for i,(du,mo) in enumerate(((30,20),(20,15),(15,10))):
    r_=23+i; put(wm,r_,1,du,blue,None,ctr); put(wm,r_,2,f"=480/A{r_}",bf,None,ctr); put(wm,r_,3,f"=B{r_}*$B$7",bf,None,ctr); put(wm,r_,4,mo,blue,None,ctr)
wm["A27"]="Dias médicos (modelo jul/2026 a jun/2027). O dia da semana de cada médico muda a cada mês."; wm["A27"].font=bold
hdr(wm,28,["Mês","Neuropediatra","Psiquiatra infantil","Pediatra"])
_bym={}
for dt,w_,di in _dd(hol): _bym.setdefault((dt.year,dt.month),{})[di]=dt
for i,((y_,m_),dd_) in enumerate(sorted(_bym.items())):
    put(wm,29+i,1,f"{m_:02d}/{y_}")
    for di in range(3): put(wm,29+i,2+di,f"{dd_[di].strftime('%d/%m')} ({WDN[dd_[di].weekday()]})")
wm["A42"]="Como funciona"; wm["A42"].font=bold
for i,t in enumerate(["O dia médico gira entre os dias da semana (cada médico muda de dia a cada mês), para que as crianças de segunda a sexta tenham acesso.",
"Cada dia médico atende só crianças que já vêm naquele dia da semana: a consulta de 30 min é encaixada antes ou depois da visita de 2 h, sem deslocamento extra.",
"Quem é chamado: crianças com menos consultas e que ainda não passaram por aquele médico; empate decide quem espera menos.",
"Todas as 16 vagas de cada dia médico ficam preenchidas (576 consultas por ano). A lista dia a dia está na aba Agenda médica.",
"Ampliar sem aumentar os dias dos médicos: consulta de 20 min leva o ciclo completo de 20 para 15 meses, e de 15 min para 10 meses. Também dá para priorizar por triagem quem precisa primeiro.",
"Faltas no dia médico liberam vaga: manter uma lista de espera de crianças que já estão na clínica naquele dia."]): wm.cell(43+i,1,t).font=bf
wm.column_dimensions["A"].width=58
for c_ in "BCD": wm.column_dimensions[c_].width=26
wm.row_dimensions[22].height=44

wa=wb.create_sheet("Agenda médica")
wa["A1"]="Agenda médica do ano-modelo: quem consulta em cada dia médico"; wa["A1"].font=title
hdr(wa,3,["Data","Dia","Médico","Horário","Criança","Visita da criança nesse dia"])
for i,(dt,doc,t,k,(d_,ini,fim)) in enumerate(sorted(_rows,key=lambda x:(x[0],x[2]))):
    r_=4+i
    for c_,v_ in enumerate((dt.strftime("%d/%m/%Y"),WDN[dt.weekday()],doc,f"{_hm(t)}–{_hm(t+30)}",k,f"{_hm(ini)}–{_hm(fim)}"),1): put(wa,r_,c_,v_)
wa.auto_filter.ref=f"A3:F{3+len(_rows)}"; wa.freeze_panes="A4"
for c_,w_ in zip("ABCDEF",(13,7,22,14,10,26)): wa.column_dimensions[c_].width=w_

# --- Comparativo
wp=wb.create_sheet("Comparativo",1)
wp["A1"]="Opções avaliadas (Opção 2 adotada em 22/10/2026)"; wp["A1"].font=title
wp["A2"]="Azul = valor informado; preto = calculado. Só a Opção 2 (adotada) tem grade nominal; as outras são comparação de capacidade. Valores da Opção 1 informados manualmente."; wp["A2"].font=bf
hdr(wp,4,["Indicador","Opção 1: 24 sessões, 3 visitas a cada 2 semanas","Opção 2 (ADOTADA): 16 sessões, 1 visita por semana","Referência: 24 sessões, 1 visita por semana de 3 h"])
rows=[("Sessões de terapia por criança a cada 4 semanas",24,16,24,"i"),
("Duração da visita (horas)",2,2,3,"i"),
("Visitas por criança a cada 4 semanas",6,4,4,"i"),
("Horas de terapia por criança a cada 4 semanas","=B5*0.5","=C5*0.5","=D5*0.5",""),
("Visitas por dia (blocos de visita)",60,60,38,"i"),
("Crianças a 100% de ocupação",200,"=C9*5","=D9*5",""),
("Blocos livres por dia no cenário de 10% de reposição",6,6,4,"i"),
("Crianças com 10% de reposição",180,"=(C9-C11)*5","=(D9-D11)*5",""),
("Crianças com horário totalmente fixo (100%)",180,"=C10","=D10",""),
("Sessões de terapia por criança por ano","=B5*Resumo!$B$5/20","=C5*Resumo!$B$5/20","=D5*Resumo!$B$5/20",""),
("Sessões de terapia por ano da clínica (100%)","=B10*B14","=C10*C14","=D10*D14",""),
("Crianças com 1ª consulta de cada médico em 1 dia/mês (192 por médico/ano)","=MIN(1,192/B10)","=MIN(1,192/C10)","=MIN(1,192/D10)",""),
("Dias de consulta por mês, por médico, para atender todas as crianças","=B10/16/12","=C10/16/12","=D10/16/12","")]
for i,(a,x,y,z,t) in enumerate(rows):
    r=5+i; put(wp,r,1,a)
    for j,v in enumerate((x,y,z)):
        c=put(wp,r,2+j,v,blue if t=="i" else bold if "Crianças" in a and "100%" in a or "10% de reposição" in a and "Crianças" in a else bf,None,ctr)
for c in "BCD":
    wp[c+"16"].number_format="0%"; wp[c+"17"].number_format="0.0"; wp[c+"14"].number_format="#,##0"; wp[c+"15"].number_format="#,##0"
wp["A19"]="Notas"; wp["A19"].font=bold
for k,t in enumerate(["As Opções 1 e 2 usam a mesma grade do dia (60 blocos); só muda quantas crianças passam por cada bloco.",
 "Opção 2 (adotada): cada bloco recebe 1 criança por dia da semana (5 por semana). Com 16 sessões por mês, cada criança tem 8 horas de terapia por mês, contra 12 horas nas outras opções.",
 "A clínica entrega o mesmo número de sessões por ano nas Opções 1 e 2; a diferença é atender mais crianças com menos sessões cada.",
 "Opção de referência (3 h por semana): 38 visitas por dia comprovadas na grade; 40 não foi concluído pelo solver. Sem grade nominal.",
 "Consultas médicas: o limite de 1 dia por mês por médico cobre bem a Opção 1, mas só cerca de 64% das crianças na Opção 2."]):
    wp.cell(20+k,1,t).font=bf
wp.column_dimensions["A"].width=64
for c in "BCD": wp.column_dimensions[c].width=30
wp.row_dimensions[4].height=48

# --- Serviços
wv=wb.create_sheet("Serviços extras")
r=1; wv.cell(r,1,"Serviços extras e de apoio").font=title; r=3
wv.cell(r,1,"Atendimento médico: ver as abas Médicos e Agenda médica.").font=bold; r+=2
wv.cell(r,1,"Odontologia: 1 dentista por dia, consultas de 40 min = 16 vagas/dia (7 de manhã, 9 à tarde)").font=bold; r+=1
hdr(wv,r,["Manhã","Tarde"]); r+=1
for i in range(9):
    put(wv,r,1,DENT[i] if i<7 else ""); put(wv,r,2,DENT[7+i]); r+=1
r+=1
wv.cell(r,1,"Psicóloga dos pais: sessões de 30 min; G = grupo (até 8 pais), I = individual").font=bold; r+=1
hdr(wv,r,["Horário","Tipo","Horário","Tipo"]); r+=1
sl=[tm(s) for s in range(22)]
for i in range(11):
    for j,k in enumerate((i,i+11)):
        t=sl[k]; put(wv,r,1+2*j,f"{t}–{tend(k)}"); put(wv,r,2+2*j,"G" if t in GRP else "I",bf,"E2EFDA" if t in GRP else None,ctr)
    r+=1
wv.cell(r,1,"Demanda estimada: 1 sessão a cada 2 visitas, cerca de 30 pais por dia com 300 crianças (27 com 270).").font=bf; r+=2
wv.cell(r,1,"Xadrez e capoeira: aulas de 1h no horário cheio, ao lado da visita (antes ou depois). Alternam a cada visita. Simulação com as 60 visitas do dia.").font=bold; r+=1
hdr(wv,r,["Aula","Xadrez (máx. 6)","Capoeira (máx. 10)"]); r+=1
for h in HRS:
    put(wv,r,1,f"{h:02d}:00–{h+1:02d}:00"); put(wv,r,2,cnt.get(("X",h),0)); put(wv,r,3,cnt.get(("C",h),0)); r+=1
put(wv,r,1,"Total/dia",bold); put(wv,r,2,f"=SUM(B{r-11}:B{r-1})",bold); put(wv,r,3,f"=SUM(C{r-11}:C{r-1})",bold); r+=1
wv.cell(r,1,"Espera média entre a visita e a aula: 4 min (máx. 30 min). Salas: 1 para xadrez, 1 para capoeira; 1 profissional por turno em cada.").font=bf; r+=2
wv.cell(r,1,"Oficinas profissionalizantes (turmas fixas): 1 sala, 7 encontros de 1h30 por dia. Distribuição PROVISÓRIA; o funcionamento das turmas ainda será definido.").font=bold; r+=1
hdr(wv,r,["Horário","Seg","Ter","Qua","Qui","Sex"]); r+=1
for j,(a,b) in enumerate(OSL):
    put(wv,r,1,f"{a}–{b}")
    for d in range(5):
        v=ogrid[j][d]; put(wv,r,2+d,v,bf,"FCE4D6" if "(mães)" in v else "DDEBF7",Alignment(wrap_text=True,vertical="center"))
    wv.row_dimensions[r].height=32; r+=1
wv.cell(r,1,"Rosa = oficinas para mães; azul = oficinas para pais.").font=bf
for c,w in zip("ABCDEF",(22,30,30,30,30,30)): wv.column_dimensions[c].width=w

# --- Equipe
we=wb.create_sheet("Equipe e turnos")
we["A1"]="Equipe e turnos (remuneração por turno: manhã 07h–12h e tarde 13h–19h)"; we["A1"].font=title
hdr(we,3,["Função","Profissionais por turno","Turnos por dia","Dias por ano","Turnos por ano"])
for i,(n,p,t) in enumerate(EQ):
    r=4+i; put(we,r,1,n); put(we,r,2,p,blue); put(we,r,3,t,blue); put(we,r,4,"=Calendário!$B$16"); put(we,r,5,f"=B{r}*C{r}*D{r}")
r=4+len(EQ)
put(we,r,1,"Médicos (3 especialidades, 1 dia/mês, 2 turnos)"); put(we,r,2,3,blue); put(we,r,3,2,blue); put(we,r,4,12,blue); put(we,r,5,f"=B{r}*C{r}*D{r}")
put(we,r+1,1,"Total de turnos por ano",bold); put(we,r+1,5,f"=SUM(E4:E{r})",bold)
we.cell(r+3,1,"Xadrez: 1 profissional por turno (2 por dia). Capoeira: 1 profissional por turno (confirmado).").font=bf
we.cell(r+4,1,"Dentista: pode ser profissionais diferentes por dia; o custo anual considera 1 por turno em todos os dias de atendimento.").font=bf
for c,w in zip("ABCDE",(48,22,16,14,16)): we.column_dimensions[c].width=w

for n in ("Ciclo Cenário 2","Ciclo Cenário 1","Crianças"): del wb[n]
wb.save(OUT+"agenda-atendimentos.xlsx")
json.dump(dict(blocks=blocks),open("data.json","w"))
print("xlsx ok")
